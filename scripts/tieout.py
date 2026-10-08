"""Tie-out: recompute every number that appears on a chart or on the page, straight from the
raw files, without importing any pipeline code, and compare with what the charts and the page
used (charts/numbers.json, data/processed/page_data.json). Writes TIEOUT.md; exits non-zero on
any mismatch.
"""
import csv, json, sys, calendar, datetime as dt
from pathlib import Path
import openpyxl

R = Path(__file__).resolve().parents[1]
RAW = R / 'data' / 'raw'
NUM = json.load(open(R / 'charts' / 'numbers.json'))
PAGE = json.load(open(R / 'data' / 'processed' / 'page_data.json'))
THIS = dt.date.today().strftime('%Y-%m')
rows, bad = [], 0


def check(section, item, recomputed, used, tol=1e-6):
    global bad
    ok = abs(recomputed - used) <= tol * max(1, abs(used))
    bad += not ok
    rows.append((section, item, recomputed, used, 'ok' if ok else 'MISMATCH'))


# ---- raw readers, written separately from process.py ----
def census_uv(flow, codes):
    """Dollars per tonne by month, summing the codes given: total value over total quantity."""
    val, qty, seen = {}, {}, set()
    for r in csv.DictReader(open(RAW / 'census' / f'{flow}.csv')):
        if r['commodity'] in codes and r['cty_name'] == 'TOTAL FOR ALL COUNTRIES':
            assert (r['commodity'], r['month']) not in seen
            seen.add((r['commodity'], r['month']))
            val[r['month']] = val.get(r['month'], 0) + float(r['value_usd'])
            qty[r['month']] = qty.get(r['month'], 0) + float(r['quantity'])
    return {m: val[m] / qty[m] for m in val if qty[m] > 0}


def fred(sid):
    return {r['observation_date'][:7]: float(r[sid]) for r in csv.DictReader(open(RAW / 'fred' / f'{sid}.csv')) if r[sid] not in ('', '.')}


def pink(name):
    ws = openpyxl.load_workbook(RAW / 'worldbank' / 'CMO-Historical-Data-Monthly.xlsx', read_only=True, data_only=True)['Monthly Prices']
    allrows = list(ws.iter_rows(values_only=True))
    j = [i for i, h in enumerate(allrows[4]) if h and str(h).strip() == name]
    assert len(j) == 1, name
    return {r[0][:4] + '-' + r[0][5:]: float(r[j[0]]) for r in allrows[6:] if r[0] and isinstance(r[j[0]], (int, float))}


def diesel_monthly():
    b = {}
    for r in csv.DictReader(open(RAW / 'eia' / 'diesel_retail_weekly.csv')):
        b.setdefault(r['period'][:7], []).append(float(r['value']))
    return {m: sum(v) / len(v) for m, v in b.items() if m < THIS}


def usgs_t1():
    """Latest workbook that carries a month wins. Returns {month: (shipments, stocks)}."""
    out = {}
    for f in sorted((RAW / 'usgs').glob('mis-*-sulfu.xlsx')):
        year = None
        for r in openpyxl.load_workbook(f, data_only=True)['T1'].iter_rows(values_only=True):
            a = str(r[0] or '').strip()
            if a.isdigit() and len(a) == 4:
                year = int(a)
            elif a in calendar.month_name[1:] and year:
                clean = lambda c: float(str(c).replace(',', '').split()[0])
                out[f'{year}-{list(calendar.month_name).index(a):02d}'] = (clean(r[4]), clean(r[5]))
    return out


sulfur, acid, dap, diesel = census_uv('exports', ('2503000010', '2503000090')), fred('WPU0613020T1'), pink('DAP'), diesel_monthly()
SER = {'sulfur_export_uv': sulfur, 'ppi_sulfuric_acid': acid, 'wb_dap': dap, 'diesel_retail': diesel,
       'ppi_diesel': fred('WPU057303'), 'wb_brent': pink('Crude oil, Brent')}
last = lambda d: max(m for m in d if m < THIS)

for key, n in NUM['run_up'].items():
    d = SER[key]; m = last(d)
    assert m == n['last_month'], (key, m, n['last_month'])
    check('Prices since January 2024', f'{key} January 2024 level', d['2024-01'], n['base'])
    check('Prices since January 2024', f'{key} {m} level', d[m], n['last'])
    check('Prices since January 2024', f'{key} percent change to {m}', 100 * (d[m] / d['2024-01'] - 1), n['pct'])
    check('Prices since January 2024', f'{key} percent change to February 2026', 100 * (d['2026-02'] / d['2024-01'] - 1), n['pct_at_war_base'])
for key, n in NUM['since_war'].items():
    d = SER[key]; m = last(d)
    assert m == n['last_month'], (key, m, n['last_month'])
    check('Prices since February 2026', f'{key} February 2026 level', d['2026-02'], n['base'])
    check('Prices since February 2026', f'{key} percent change to {m}', 100 * (d[m] / d['2026-02'] - 1), n['pct'])

# world production, from the hand-keyed table (the independent tie-out re-reads the PDF)
w = [r for r in csv.DictReader(l for l in open(R / 'data' / 'manual' / 'usgs_mcs2026_world_production.csv') if not l.startswith('#'))]
gulf = sum(float(r['y2025']) for r in w if r['country'] in ('Iran', 'Kuwait', 'Qatar', 'Saudi Arabia', 'United Arab Emirates'))
world = float(next(r['y2025'] for r in w if r['country'].startswith('World')))
check('World production 2025', 'Five Gulf producers, thousand t', gulf, NUM['world']['gulf_kt'])
check('World production 2025', 'World total, thousand t', world, NUM['world']['world_kt'])
check('World production 2025', 'Gulf share, percent', 100 * gulf / world, NUM['world']['gulf_pct'])
# USGS prints a world total a little above the sum of its own rows (130 in 2025, 650 in 2024; see DATASETS.md), so allow 1%.
check('World production 2025', 'Sum of the rows USGS lists vs its printed world total, thousand t (within 1%)',
      sum(float(r['y2025']) for r in w if not r['country'].startswith('World')), world, tol=0.01)
for r in w:
    if r['country'] in NUM['world']['countries']:
        check('World production 2025', r['country'], float(r['y2025']), NUM['world']['countries'][r['country']])

# days of supply
t1 = usgs_t1(); m = max(t1); n = NUM['days']
assert m == n['sulfur_last_month']
days = {k: st / (sh / calendar.monthrange(int(k[:4]), int(k[5:]))[1]) for k, (sh, st) in t1.items()}
check('Days of supply', f'Sulfur producer stocks {m}, thousand t', t1[m][1], n['sulfur_stocks_kt'])
check('Days of supply', f'Sulfur shipments {m}, thousand t', t1[m][0], n['sulfur_shipments_kt'])
check('Days of supply', f'Sulfur days {m}', days[m], n['sulfur_days'])
check('Days of supply', 'Sulfur days, lowest month', min(days.values()), n['sulfur_days_min'])
check('Days of supply', 'Sulfur days, highest month', max(days.values()), n['sulfur_days_max'])
dd = {r['period']: float(r['value']) for r in csv.DictReader(open(RAW / 'eia' / 'distillate_days_weekly.csv')) if r['period'] >= '2024-01-01'}
check('Days of supply', f'Distillate days, week of {max(dd)}', dd[max(dd)], n['distillate_days'])
check('Days of supply', 'Distillate days, lowest week since 2024', min(dd.values()), n['distillate_days_min'])
check('Days of supply', 'Distillate days, highest week since 2024', max(dd.values()), n['distillate_days_max'])

# sulfur cost inside a ton of DAP
n = NUM['dap_cost']; k = n['s_per_dap']; dm = n['last_month']
assert dm == min(last(sulfur), last(dap))
for label, mm, a, b, c in [('first', '2024-01', 'sulfur_cost_first', 'dap_first', 'share_first_pct'),
                           ('February 2026', '2026-02', 'sulfur_cost_war_base', 'dap_war_base', 'share_war_base_pct'),
                           ('last', dm, 'sulfur_cost_last', 'dap_last', 'share_last_pct')]:
    check('Sulfur in a ton of DAP', f'Sulfur cost {mm}, $ per t DAP', k * sulfur[mm], n[a])
    check('Sulfur in a ton of DAP', f'DAP price {mm}, $ per t', dap[mm], n[b])
    check('Sulfur in a ton of DAP', f'Sulfur share of DAP price {mm}, percent', 100 * k * sulfur[mm] / dap[mm], n[c])

# posted prices (hand-compiled table; the independent tie-out re-opens the sources)
ad = {r['month']: float(r['price']) for r in csv.DictReader(l for l in open(R / 'data' / 'manual' / 'adnoc_osp.csv') if not l.startswith('#'))}
n = NUM['posted']; lastm = max(ad); peak = max(ad.values())
check('Posted prices', 'Adnoc February 2026, $ per t', ad['2026-02'], n['adnoc_war_base'])
check('Posted prices', 'Adnoc peak, $ per t', peak, n['adnoc_peak'])
check('Posted prices', f'Adnoc {lastm}, $ per t', ad[lastm], n['adnoc_last'])
check('Posted prices', 'Adnoc percent change, February 2026 to peak', 100 * (peak / ad['2026-02'] - 1), n['pct_war_base_to_peak'])
check('Posted prices', f'Adnoc percent change, February 2026 to {lastm}', 100 * (ad[lastm] / ad['2026-02'] - 1), n['pct_war_base_to_last'])
check('Posted prices', f'Adnoc percent change, peak to {lastm}', 100 * (ad[lastm] / peak - 1), n['pct_peak_to_last'])

# stockpiles
cn = [(r['date'], float(r['million_t'])) for r in csv.DictReader(l for l in open(R / 'data' / 'manual' / 'china_port_stocks.csv') if not l.startswith('#'))]
n = NUM['stockpiles']; low = min(v for _, v in cn)
check('Stockpiles', 'China ports, first reading, million t', cn[0][1], n['china_first'])
check('Stockpiles', 'China ports, lowest reading, million t', low, n['china_low'])
check('Stockpiles', 'China ports, latest reading, million t', cn[-1][1], n['china_last'])
check('Stockpiles', 'China ports, percent change first to lowest', 100 * (low / cn[0][1] - 1), n['china_pct_first_to_low'])
ab = {}
for f in sorted((RAW / 'aer').glob('st3_sulphur_*.xlsx')):
    for r in openpyxl.load_workbook(f, data_only=True)['Data'].iter_rows(values_only=True):
        if r[1] and str(r[1]).strip() == 'Closing Inventory':
            for i, v in enumerate(r[3:15]):
                if v:
                    ab[f'{f.stem[-4:]}-{i + 1:02d}'] = float(v)
am = max(ab)
assert am == n['alberta_last_month']
check('Stockpiles', 'Alberta closing inventory February 2026, million t', ab['2026-02'] / 1e6, n['alberta_war_base'])
check('Stockpiles', f'Alberta closing inventory {am}, million t', ab[am] / 1e6, n['alberta_last'])
check('Stockpiles', f'Alberta change February 2026 to {am}, t', ab[am] - ab['2026-02'], n['alberta_change_since_war_base_t'])
check('Stockpiles', f'Alberta percent change February 2026 to {am}', 100 * (ab[am] / ab['2026-02'] - 1), n['alberta_pct_since_war_base'])
check('Stockpiles', f'U.S. producer stocks {m}, thousand t', t1[m][1], n['us_last_kt'])

# the page: every point of every series (the page stores four decimals)
for key, src in [('sulfur', sulfur), ('acid', acid), ('dap', dap), ('diesel', diesel)]:
    pts = dict(PAGE['series'][key])
    want = {mm: v for mm, v in src.items() if '2024-01' <= mm < THIS}
    assert set(pts) == set(want), f'page months differ for {key}'
    worst = max(abs(pts[mm] - want[mm]) for mm in want)
    check('Interactive page', f'{key}: {len(want)} months, largest difference from raw', worst, 0, tol=6e-5)

out = ['# Tie-out', '', f'Run {dt.datetime.now():%B %d, %Y %H:%M}. Every number below is recomputed from the raw files in `data/raw/` by',
       '`scripts/tieout.py`, which shares no code with the pipeline, and compared with the number the chart or page used.', '',
       f'**{len(rows) - bad} of {len(rows)} match.**' + (f' **{bad} MISMATCH.**' if bad else ''), '']
sec = None
for s, item, a, b, st in rows:
    if s != sec:
        out += ['', f'## {s}', '', '| Item | Recomputed from raw | Used on chart or page | |', '|---|---:|---:|---|']; sec = s
    out.append(f'| {item} | {a:,.4f} | {b:,.4f} | {st} |')
out += ['', '## Not covered by this script', '',
        '- The Adnoc price table and the China port stock table are hand-compiled from trade press reports. This script reads the compiled files; the sources are listed in `research/BRIEF.md`.\n- The two USGS annual tables are hand-keyed from the PDF. This script reads the hand-keyed file; the independent tie-out re-reads the PDF.',
        '- The 0.4 metric tons of sulfur per metric ton of DAP is an input, not a measurement. Its source is in `research/BRIEF.md`.',
        '- Whether a source itself is right (for example the Census export unit value as a stand-in for a sulfur price) is covered in `DATASETS.md`, not here.']
(R / 'TIEOUT.md').write_text('\n'.join(out) + '\n')
print(f'{len(rows) - bad} of {len(rows)} match' + (f', {bad} MISMATCH' if bad else ''))
sys.exit(1 if bad else 0)
