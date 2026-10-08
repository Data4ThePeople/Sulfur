"""Turn the raw pulls into one file the charts, the page and the tie-out all read.

Output: data/processed/series.json. Nothing here is typed in by hand; the two
hand-keyed USGS tables in data/manual/ are read as files.
"""
import csv, json, re, calendar, datetime as dt
from collections import defaultdict
import openpyxl
from common import RAW, ROOT, PROCESSED, WAR_START, write_json

MONTHS = {m: i for i, m in enumerate(calendar.month_name) if m}


def read_csv(path):
    with open(path, encoding='utf-8-sig') as f:
        return list(csv.DictReader(r for r in f if not r.startswith('#')))


def unique(pairs, what):
    ks = [k for k, _ in pairs]
    assert len(ks) == len(set(ks)), f'duplicate periods in {what}'
    return sorted(pairs)


def eia(name):
    return unique([(r['period'], float(r['value'])) for r in read_csv(RAW / 'eia' / f'{name}.csv')], name)


def fred(sid):
    rows = read_csv(RAW / 'fred' / f'{sid}.csv')
    return unique([(r['observation_date'][:7], float(r[sid])) for r in rows if r[sid] not in ('', '.')], sid)


def pink_sheet():
    ws = openpyxl.load_workbook(RAW / 'worldbank' / 'CMO-Historical-Data-Monthly.xlsx', read_only=True, data_only=True)['Monthly Prices']
    rows = list(ws.iter_rows(values_only=True))
    updated = rows[3][0]
    names, units = rows[4], rows[5]
    want = {'Phosphate rock': 'rock', 'DAP': 'dap', 'TSP': 'tsp', 'Urea': 'urea', 'Crude oil, Brent': 'brent'}
    cols = {}
    for j, n in enumerate(names):
        k = want.get(str(n).strip()) if n else None
        if k:
            assert k not in cols, f'two Pink Sheet columns named {n}'
            cols[k] = j
    assert set(cols) == set(want.values()), f'Pink Sheet columns missing: {set(want.values()) - set(cols)}'
    out = {k: [] for k in cols}
    for r in rows[6:]:
        if not r[0] or not re.fullmatch(r'\d{4}M\d{2}', str(r[0])):
            continue
        m = r[0].replace('M', '-')
        for k, j in cols.items():
            if isinstance(r[j], (int, float)):
                out[k].append((m, float(r[j])))
    return {k: unique(v, 'pink ' + k) for k, v in out.items()}, updated, {k: units[j] for k, j in cols.items()}


def num(c):
    """A USGS cell: number, or a string like '642 r' or '7,370 r'. Returns None for W, NA, <0.5."""
    if isinstance(c, (int, float)):
        return float(c)
    m = re.fullmatch(r'\s*([\d,]+(?:\.\d+)?)\s*[a-z]?\s*', str(c or ''))
    return float(m.group(1).replace(',', '')) if m else None


def usgs_table(ws, ncols):
    """Month rows of a USGS monthly table, keyed YYYY-MM. Skips year totals and year-to-date rows."""
    out, year = {}, None
    for r in ws.iter_rows(values_only=True):
        a = str(r[0] or '').strip()
        if re.fullmatch(r'\d{4}', a):
            year = int(a)
        elif a in MONTHS and year:
            vals = [num(c) for c in r[1:] if c is not None][:ncols]
            assert len(vals) == ncols, f'short row {a} {year} in {ws.title}'
            out[f'{year}-{MONTHS[a]:02d}'] = vals
    return out


def usgs():
    """Each monthly workbook restates earlier months, sometimes revised. The latest workbook
    that carries a month wins, and first-published values are kept so revisions can be measured."""
    files = sorted((RAW / 'usgs').glob('mis-*-sulfu.xlsx'))
    t1, first, t2, t4 = {}, {}, {}, {}
    for f in files:   # oldest first, so later files overwrite
        wb = openpyxl.load_workbook(f, data_only=True)
        for m, v in usgs_table(wb['T1'], 5).items():
            first.setdefault(m, v)
            t1[m] = v
        t2.update(usgs_table(wb['T2'], 4))
        t4.update(usgs_table(wb['T4'], 5))
    cols = ['prod_petroleum', 'prod_gas', 'prod_total', 'shipments', 'stocks']
    series = {c: [(m, t1[m][i]) for m in sorted(t1)] for i, c in enumerate(cols)}
    series['stocks_first_published'] = [(m, first[m][4]) for m in sorted(first)]
    # Days of stock: end-of-month producer stocks over that month's daily shipments.
    series['stock_days'] = [(m, t1[m][4] / (t1[m][3] / calendar.monthrange(int(m[:4]), int(m[5:]))[1])) for m in sorted(t1)]
    series['import_value_per_t'] = [(m, t2[m][3] / t2[m][2]) for m in sorted(t2)]   # $000 over 000 t is $/t
    series['export_value_per_t'] = [(m, t4[m][1] / t4[m][0]) for m in sorted(t4)]
    return series, files[-1].name


def census(flow):
    """Monthly all-country totals per 10-digit code: value, tonnes, and dollars per tonne."""
    tot = {}
    for r in read_csv(RAW / 'census' / f'{flow}.csv'):
        if r['cty_code'] == '-':
            k = (r['commodity'], r['month'])
            assert k not in tot, f'duplicate census total {k}'
            tot[k] = (float(r['value_usd']), float(r['quantity']))
    out = defaultdict(list)
    for (code, m), (v, q) in sorted(tot.items()):
        out[code].append((m, v, q, v / q if q else None))
    return out


def usda():
    rows = json.load(open(RAW / 'usda' / 'ams_3195_prices.json'))
    want = {'DAP (Diammonium Phosphate 18-46-0)': 'dap', 'MAP (Monoammonium Phosphate 11-52-0)': 'map',
            'No. 2 Diesel (Farm)': 'farm_diesel', 'Urea (46-0-0)': 'urea', 'Potash (White 0-0-62)': 'potash'}
    out = defaultdict(list)
    for r in rows:
        k = want.get(r['class'])
        if k and r['price_avg'] is not None:
            d = dt.datetime.strptime(r['report_begin_date'], '%m/%d/%Y').strftime('%Y-%m-%d')
            out[k].append((d, float(r['price_avg'])))
    return {k: unique(v, 'usda ' + k) for k, v in out.items()}


def statcan():
    rows = read_csv(RAW / 'statcan' / '25100036.csv')
    return unique([(r['REF_DATE'], float(r['VALUE']) / 1000) for r in rows
                   if r['GEO'] == 'Canada' and r['Supply of processing plant products'] == 'Sulphur' and r['VALUE']], 'statcan')


def monthly_mean(weekly):
    b = defaultdict(list)
    for d, v in weekly:
        b[d[:7]].append(v)
    return [(m, sum(v) / len(v)) for m, v in sorted(b.items())]


if __name__ == '__main__':
    pink, pink_updated, pink_units = pink_sheet()
    usgs_m, usgs_latest = usgs()
    imp, exp = census('imports'), census('exports')
    diesel_w = eia('diesel_retail_weekly')
    world = read_csv(ROOT / 'data' / 'manual' / 'usgs_mcs2026_world_production.csv')
    us = {r['item']: {k: float(v) for k, v in r.items() if k != 'item'} for r in read_csv(ROOT / 'data' / 'manual' / 'usgs_mcs2026_us_salient.csv')}
    out = {
        'war_start': WAR_START,
        'pulled': {p.parent.name: p.read_text().strip() for p in sorted(RAW.glob('*/pulled.txt'))},
        'pink_updated': pink_updated, 'usgs_latest_file': usgs_latest,
        'weekly': {
            'diesel_retail': diesel_w,                          # $/gal
            'distillate_stocks': eia('distillate_stocks_weekly'),   # thousand barrels
            'distillate_days': eia('distillate_days_weekly'),
        },
        'biweekly_illinois': usda(),                            # $/short ton; farm diesel $/gal
        'monthly': {
            'diesel_retail': monthly_mean(diesel_w),            # mean of the weeks that fall in the month
            'ppi_sulfuric_acid': fred('WPU0613020T1'),
            'ppi_phosphates': fred('WPU065202'),
            'ppi_phosphatic_fertilizer_mfg': fred('PCU325312325312'),
            'ppi_diesel': fred('WPU057303'),
            'wb_dap': pink['dap'], 'wb_tsp': pink['tsp'], 'wb_rock': pink['rock'], 'wb_urea': pink['urea'], 'wb_brent': pink['brent'],
            # Census unit values, $/t: crude or unrefined sulfur (2503.00.0010), sulfuric acid, DAP
            'sulfur_export_uv': [(m, uv) for m, v, q, uv in exp['2503000010'] if uv],
            'sulfur_import_uv': [(m, uv) for m, v, q, uv in imp['2503000010'] if uv],
            'sulfur_export_t': [(m, q) for m, v, q, uv in exp['2503000010']],
            'sulfur_import_t': [(m, q) for m, v, q, uv in imp['2503000010']],
            'acid_import_uv': [(m, uv) for m, v, q, uv in imp['2807000000'] if uv],
            'acid_import_t': [(m, q) for m, v, q, uv in imp['2807000000']],
            'dap_export_uv': [(m, uv) for m, v, q, uv in exp['3105300000'] if uv],
            'canada_gas_plant_sulphur_kt': statcan(),
            **{'usgs_' + k: v for k, v in usgs_m.items()},
        },
        'world_production_kt': [{'country': r['country'], 'y2024': float(r['y2024']), 'y2025': float(r['y2025']), 'gulf': r['gulf'] == '1'} for r in world],
        'us_salient': us,
    }
    PROCESSED.mkdir(parents=True, exist_ok=True)
    write_json(PROCESSED / 'series.json', out)
    for k, v in out['monthly'].items():
        print(f'{k:36s} {len(v):4d} {v[0][0]} to {v[-1][0]}  last {v[-1][1]:.2f}')
